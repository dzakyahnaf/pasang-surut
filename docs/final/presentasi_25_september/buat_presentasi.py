"""PPT final 25 September: objek PowerPoint, bukti repo, catatan dan navigasi.

Dependency khusus materi mengikuti ../presentasi/requirements-materi.txt.
Tidak mengubah aplikasi, potret, atau database.
"""
from __future__ import annotations
import csv
import hashlib
import json
import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT/'.deploy-local/presentation-tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
import qrcode

ASSET = OUT/'aset'; ASSET.mkdir(exist_ok=True)
OLD = OUT.parent/'presentasi/aset'
NAME = 'Pasang_Surut_Final_2026_25Sep'
APP = 'https://pasang-surut.vercel.app/app'
REPO = 'https://github.com/dzakyahnaf/pasang-surut'
LOCAL = 'http://127.0.0.1:5182'
REPORT = REPO+'/blob/main/docs/final/audit_asal_data_25_september.md'
TIDE_SOURCE = 'https://www.ioc-sealevelmonitoring.org/station.php?code=sema'
PUPR = 'https://sda.pu.go.id/balai/bbwspemalijuana/pages/posts/perkuat-pengendalian-rob-pantura-wapres-tinjau-pembangunan-tol-semrang-demak-seksi-1'
DATA = json.loads((OLD/'skenario.json').read_text(encoding='utf-8'))
TIDE = json.loads((OUT.parent/'data/uji_pasut_independen_25_september.json').read_text())
AUDIT = json.loads((OUT.parent/'bukti/audit_regresi_22_verifikasi.json').read_text())
NAVY='0B1F2A'; PANEL='12303E'; BG='F2F5F6'; PALE='E3EAEC'; LINE='CBD7DB'
INK='0B1F2A'; MUTED='46626F'; TEAL='1C7F9E'; CYAN='4FB3C4'; AMBER='FFB020'
WHITE='FFFFFF'; LIGHT='E8F1F4'; SOFT='ABC0CA'; RED='C8322B'
FONT='Arial'; DISPLAY='Arial Narrow'
P=Presentation();P.slide_width=Inches(13.333333);P.slide_height=Inches(7.5)
P.core_properties.title='PASANG SURUT | Final DSDC ANFORCOM 2026'
P.core_properties.author='trio la albiceleste | ITS Surabaya'
P.core_properties.subject='Perencanaan rute dan waktu di pesisir Semarang'
P.core_properties.comments='Revisi 25 September 2026. 9 slide utama, 9 lampiran. Bukti dan batas mengikuti audit data 25 September.'
notes=[]; links=[]; slides=[]

def dec(n,d=1): return f'{n:.{d}f}'.replace('.',',')
def rgb(v):return RGBColor.from_string(v)
def box(s,x,y,w,h,color,stroke=None):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid();a.fill.fore_color.rgb=rgb(color)
    if stroke:a.line.color.rgb=rgb(stroke);a.line.width=Pt(.7)
    else:a.line.fill.background()
    return a
def text(s,value,x,y,w,h,size=22,color=INK,bold=False,font=FONT,align=None,url=None):
    a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=a.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,v in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_before=Pt(0);p.space_after=Pt(4);p.line_spacing=1.05
        if align is not None:p.alignment=align
        r=p.add_run();r.text=v;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
    if url:a.click_action.hyperlink.address=url
    return a
def rule(s,x,y,w,color=LINE):return box(s,x,y,w,.016,color)
def tag(s,value,x,y,w=10,color=TEAL):return text(s,value.upper(),x,y,w,.3,11.5,color,True)
def footer(s,n,source='',dark=False,url=None):
    rule(s,.55,7.06,12.23,PANEL if dark else LINE)
    text(s,source or 'trio la albiceleste | ITS Surabaya | ANFORCOM 2026',.55,7.15,11.6,.22,9.5,SOFT if dark else MUTED,url=url)
    text(s,f'{n:02}',12.36,7.12,.4,.28,12,SOFT if dark else MUTED,align=PP_ALIGN.RIGHT)
def slide(n,chapter,title,sub='',dark=False,source='',url=None):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY if dark else BG)
    slides.append(s);tag(s,chapter,.58,.35,color=CYAN if dark else TEAL)
    text(s,title,.56,.86,12.2,1.14 if '\n' in title else .66,34,LIGHT if dark else INK,True,DISPLAY)
    if sub:text(s,sub,.59,1.64 if '\n' not in title else 2.08,12.05,.5,17,SOFT if dark else MUTED)
    footer(s,n,source,dark,url)
    return s
def note(s,title,seconds,speaker,body,cues='',sources=''):
    start=sum(n['seconds'] for n in notes);end=start+seconds
    time=f'{start//60:02}:{start%60:02} sampai {end//60:02}:{end%60:02}' if seconds else 'Lampiran untuk tanya jawab'
    n={'slide':len(slides),'title':title,'seconds':seconds,'time':time,'speaker':speaker,'body':body,'cues':cues,'sources':sources}
    notes.append(n)
    s.notes_slide.notes_text_frame.text=f'{title}\nWAKTU: {time}\nPEMBICARA: {speaker} (pembagian latihan)\n\n{body}\n\nAKSI: {cues}\n\nSUMBER: {sources}'
def picture(s,path,x,y,w,h=None):
    return s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h) if h else None)
def button(s,value,x,y,w,target=None,url=None,dark=False,size=11,height=.42):
    a=box(s,x,y,w,height,PANEL if dark else PALE)
    b=text(s,value,x+.1,y+.09,w-.2,height-.14,size,LIGHT if dark else TEAL,True,url=url)
    if target:links.extend([(a,target),(b,target)])
    if url:a.click_action.hyperlink.address=url
def circle(s,x,y,size,color):
    a=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(size),Inches(size));a.fill.solid();a.fill.fore_color.rgb=rgb(color);a.line.fill.background();return a
def paths(s,lines,color,width):
    b=None
    for line in lines:
        if len(line)<2:continue
        pts=[(round(x*1000),round(y*1000)) for x,y in line]
        if b is None:b=s.shapes.build_freeform(*pts[0],scale=Inches(1)/1000)
        else:b.move_to(*pts[0])
        b.add_line_segments(pts[1:],close=False)
    if b:
        a=b.convert_to_shape();a.fill.background();a.line.color.rgb=rgb(color);a.line.width=Pt(width);return a
ROADS=json.loads((ROOT/'data/processed/ruas_jalan.geojson').read_text())
COAST=json.loads((ROOT/'data/processed/garis_pantai.geojson').read_text())
def map_panel(s,x,y,w,h,route=None,dark=False):
    # Koordinat geografis memakai skala setara, tidak meregangkan garis jalan.
    cx,cy=110.447,-6.960
    dx=.092;dy=dx*math.cos(math.radians(cy))*h/w
    bounds=(cx-dx/2,cy-dy/2,cx+dx/2,cy+dy/2)
    def inside(p):return bounds[0]<=p[0]<=bounds[2] and bounds[1]<=p[1]<=bounds[3]
    def xy(p):return x+(p[0]-bounds[0])/(bounds[2]-bounds[0])*w,y+(bounds[3]-p[1])/(bounds[3]-bounds[1])*h
    box(s,x,y,w,h,PANEL if dark else PALE)
    def clipped(seqs):
        lines=[]
        for seq in seqs:
            part=[]
            for p in seq:
                if inside(p):part.append(xy(p))
                else:
                    if len(part)>1:lines.append(part)
                    part=[]
            if len(part)>1:lines.append(part)
        return lines
    roadseq=[f['geometry']['coordinates'] for f in ROADS['features'] if f['properties'].get('jenis') in ['trunk','primary','secondary','tertiary','trunk_link','primary_link']]
    paths(s,clipped(roadseq),'345565' if dark else 'A0B5BF',.65)
    co=[]
    for f in COAST['features']:
        co.extend([f['geometry']['coordinates']] if f['geometry']['type']=='LineString' else f['geometry']['coordinates'])
    paths(s,clipped(co),CYAN if dark else TEAL,1.3)
    if route:
        for f in DATA[route]['rute']['features']:
            if f.get('geometry') and f['properties']['jenis']=='rute_sadar_rob':
                paths(s,clipped([f['geometry']['coordinates']]),INK,4.5)
                paths(s,clipped([f['geometry']['coordinates']]),AMBER,2.6)
    text(s,'LAUT JAWA',x+.18,y+.16,2.2,.25,11.5,CYAN if dark else TEAL,True)
    for coord,label,ox,oy,ww in [([110.42744,-6.96374],'TAWANG',-.55,.14,1.4),([110.46934,-6.94825],'TERBOYO',-.56,-.38,1.5),([110.41853,-6.94959],'TANJUNG EMAS',-.25,-.40,2)]:
        px,py=xy(coord);circle(s,px-.06,py-.06,.12,AMBER if dark else INK)
        text(s,label,px+ox,py+oy,ww,.27,12,LIGHT if dark else INK,True)
    text(s,'SEMARANG',x+w-2.38,y+h-.38,2.2,.28,13,CYAN if dark else TEAL,True)
    text(s,'U ↑',x+w-.48,y+.16,.4,.3,12,LIGHT if dark else INK)
def grid_table(s,headers,rows,x,y,widths,rowh=.75,head=.53,size=19):
    full=sum(widths);box(s,x,y,full,head,PANEL)
    for j,(v,w) in enumerate(zip(headers,widths)):
        text(s,v,x+sum(widths[:j])+.13,y+.12,w-.26,head-.17,13,LIGHT,True)
    for i,row in enumerate(rows):
        yy=y+head+i*rowh
        if i%2==0:box(s,x,yy,full,rowh,PALE)
        for j,(v,w) in enumerate(zip(row,widths)):
            text(s,str(v),x+sum(widths[:j])+.13,yy+.16,w-.26,rowh-.20,size,INK,j==0)
    return y+head+len(rows)*rowh
def stat(s,value,caption,x,y,w=3.6,dark=False):
    text(s,value,x,y,w,.69,36,CYAN if dark else TEAL,True,DISPLAY)
    text(s,caption,x,y+.82,w,.88,18,LIGHT if dark else MUTED)

# 1. Produk, wilayah, dan satu pertanyaan.
s=slide(1,'Final DSDC | ANFORCOM 2026','',dark=True,
        source='Peta konteks dari OpenStreetMap. Garis ambar: contoh rute model 26 September 08.00 WIB.',url='https://www.openstreetmap.org/copyright')
text(s,'PASANG\nSURUT',.6,1.06,5.55,2.18,61,LIGHT,True,DISPLAY)
text(s,'Berangkat kapan,\nlewat mana?',.64,3.28,5.1,1.27,33,LIGHT,True,DISPLAY)
text(s,'Perencanaan perjalanan\ndi pesisir Semarang.',.64,4.85,5.2,.9,23,SOFT)
map_panel(s,6.26,1.02,6.5,4.62,route='8',dark=True)
tag(s,'Prototipe berbasis indeks kerentanan rob',6.4,5.95,6.1,CYAN)
text(s,'trio la albiceleste | ITS Surabaya',.64,6.0,5.3,.35,16,LIGHT,True)
text(s,"Muhammad Dzaky Ahnaf · Daffa Rajendra Priyatama\nNaufal Syafi' Hakim",.64,6.46,10.8,.49,12,SOFT)
note(s,'Pembuka',30,'Dzaky',
 'Selamat pagi, Bapak dan Ibu juri. Kami tim trio la albiceleste dari ITS. PASANG SURUT membantu pengendara membandingkan rute dan jam berangkat di pesisir Semarang. Hari ini kami memakai satu perjalanan, dari akses Stasiun Tawang menuju kawasan industri Terboyo. Dari perjalanan itu, kami akan memperlihatkan apa yang sudah bekerja dan batas data yang masih perlu kami perbaiki.',
 'Tunjuk Tawang dan Terboyo. Jangan membuka dengan daftar teknologi.', 'README tim; data OSM; potret demo 26 September.')

# 2. Kebutuhan yang konkret, tanpa persona atau wawancara fiktif.
s=slide(2,'Persoalan pengguna','Dari Tawang ke Terboyo, apa yang perlu diputuskan?',
 'Ilustrasi kebutuhan pengendara motor menuju kawasan industri. Belum hasil wawancara pengguna.',
 source='Konteks penanganan rob: BBWS Pemali Juana, 15 Februari 2026. Lokasi: OpenStreetMap.',url=PUPR)
map_panel(s,.6,2.34,6.4,3.70)
tag(s,'Perjalanan yang sama',7.42,2.4,4.9)
text(s,'Harus berangkat pagi?',7.42,2.95,5.15,.55,27,INK,True,DISPLAY)
text(s,'Bandingkan ruas yang dilalui\ndan tambahan waktu memutar.',7.42,3.64,5.14,.92,23,MUTED)
rule(s,7.42,4.80,5.1)
text(s,'Jamnya bisa digeser?',7.42,5.07,5.15,.55,27,INK,True,DISPLAY)
text(s,'Bandingkan pilihan rute pada jam lain.',7.42,5.64,5.1,.75,22,MUTED)
text(s,'Rob berkaitan dengan pasang laut. Kondisi tiap jalan juga dipengaruhi lingkungan setempat.',.63,6.63,12.03,.32,15,MUTED)
note(s,'Kebutuhan perjalanan',50,'Dzaky',
 'Tawang dan Terboyo berada di kawasan pesisir yang menjadi konteks proyek ini. Rob di Semarang juga masih masuk agenda penanganan pemerintah pada 2026. Kami mengambil kebutuhan yang sederhana: seseorang perlu menuju kawasan industri dengan motor. Kalau jam berangkatnya tetap, ia perlu memahami rute dan konsekuensi memutar. Kalau jadwalnya bisa digeser, ia ingin membandingkan jam lain. Ilustrasi ini belum kami uji lewat wawancara. Karena itu, kebutuhan dan kemudahan pemakaiannya tetap masuk rencana uji pengguna setelah final. Daffa akan memperlihatkan alur yang sudah bisa dicoba.',
 'Daffa menyiapkan tab lokal saat Dzaky berbicara. Perpindahan pembicara dilakukan setelah kalimat terakhir.', 'BBWS Pemali Juana 15 Februari 2026; user mengonfirmasi belum ada uji pengguna pada 25 September.')

# 3. Aplikasi asli menjadi orientasi sebelum live demo.
s=slide(3,'Alur aplikasi','Pilih perjalanan, lalu geser jamnya.',
 'Asal, tujuan, dan moda tetap terlihat ketika pengguna membandingkan waktu.',
 source='Tangkapan aplikasi asli, potret 26 September 2026. Detail hasil diperbesar saat demo.',url=APP)
picture(s,OLD/'demo_08.png',.6,2.23,8.02)
for i,(title,body) in enumerate([('Pilih akses tujuan','Tawang → Terboyo\nModa motor'),('Geser Pita Pasut','Coba 08.00, lalu 13.00\npada tanggal yang sama.'),('Baca hasilnya','Rute, waktu, paparan,\ndan biaya memutar.')]):
    y=2.42+i*1.42
    circle(s,8.96,y,.37,TEAL);text(s,str(i+1),9.03,y+.05,.22,.22,12,WHITE,True)
    text(s,title,9.49,y,3.22,.44,23,INK,True,DISPLAY)
    text(s,body,9.49,y+.57,3.2,.82,18,MUTED)
note(s,'Alur pemakaian',45,'Daffa',
 'Di layar ini, pengguna memilih titik berangkat, tujuan, dan kendaraan. Empat akses tempat sudah tersedia sebagai pilihan cepat. Pita Pasut di bawah mengatur jam yang dibandingkan. Setelah jam berubah, peta dan hasil rute ikut dihitung. Panel kiri memperlihatkan waktu tempuh dan peringatan; selisih jarak serta bahan bakar bisa dibaca pada panel biaya. Nama tempat membantu orientasi, tetapi titik aksesnya masih hasil pelekatan ke jaringan jalan. Sekarang kita coba perjalanan yang sama.',
 'Tunjuk lokasi kontrol, pita waktu, dan panel hasil. Tidak perlu membacakan seluruh teks pada screenshot.', 'docs/final/presentasi/aset/demo_08.png; tujuan_cepat.geojson.')

# 4. Satu slide tetap terbuka selama perpindahan ke aplikasi.
s=slide(4,'Demo langsung | 2 menit 25 detik','Satu perjalanan. Kita coba dua jam berangkat.',
 'Motor · akses Tawang → akses Terboyo · Sabtu, 26 September 2026',dark=True,
 source='Skenario potret bertanggal. Angka berasal dari API lokal; waktu tempuh dan genangan adalah hasil model.')
for x,hour,minutes,km,wet,detail in [(0.65,'08.00','13,9','10,58','14','Masih melewati ruas dengan estimasi genangan.'),(7.10,'13.00','7,0','6,32','0','Rute model sama dengan rute pembanding.')]:
    tag(s,hour+' WIB',x,2.44,4,CYAN)
    text(s,minutes+' menit',x,3.02,5.15,.82,43,LIGHT,True,DISPLAY)
    text(s,km+' km  ·  '+wet+' ruas basah dalam model',x,4.08,5.47,.85,22,LIGHT)
    text(s,detail,x,5.16,5.44,.86,21,SOFT)
box(s,6.43,2.4,.02,3.60,'1B4356')
button(s,'Buka demo lokal',.66,6.36,2.16,url=LOCAL,dark=True)
button(s,'Cadangan 08.00',3.03,6.36,2.16,target=11,dark=True)
button(s,'Cadangan 13.00',5.40,6.36,2.16,target=12,dark=True)
text(s,'0 ruas model ≠ bukti jalan kering',8.0,6.45,4.72,.28,16,AMBER,True)
note(s,'Demo dua jam',145,'Daffa',
 'Kami memakai potret bertanggal 26 September agar hasil demo dapat diulang. Asal dan tujuannya tetap, dengan moda motor. Pada pukul delapan, rute sadar rob menghasilkan 13,9 menit untuk 10,58 kilometer. Peringatannya masih muncul karena rute ini tetap melewati ruas yang diberi genangan oleh model. Jadi, jangan membacanya sebagai jaminan aman.\n\nSekarang jamnya digeser ke pukul satu siang. Pada skenario ini hasilnya menjadi 7 menit dan 6,32 kilometer. Rutenya sama dengan pembanding, dan model tidak menandai ruas basah pada rute itu. Kondisi lapangan tetap perlu diperiksa. Kalau jadwal pengguna fleksibel, perbandingan seperti ini bisa menjadi bahan memilih jam. Kalau harus berangkat pagi, biaya memutarnya tetap terlihat. Dzaky akan menjelaskan bagaimana angka tadi dibentuk.',
 '00-20 detik: Alt+Tab ke lokal yang siap. 20-65: Tawang, Terboyo, motor, jam 08.00 dan peringatan. 65-105: geser ke 13.00. 105-130: bandingkan pilihan. 130-145: kembali ke slide 5. Bila satu permintaan tertahan 10 detik, langsung buka slide 11 lalu 12; jangan mengulang-ulang request.',
 'aset/skenario.json pada paket 23 September; 919aa345-7492-4a53-9e7c-c34086775d18.')

# 5. Diagram metode yang membedakan sumber, indeks, asumsi, dan rute.
s=slide(5,'Cara kerja','Data wilayah dan pasut membentuk skenario ruas.',
 'Produksi memakai indeks kerentanan berbasis aturan.',
 source='Kode aktif: kerentanan.py, pasut.py, genangan.py, routing.py. Audit sumber 25 September 2026.',url=REPORT)
for x,title,body,col in [(.65,'Kondisi wilayah','Elevasi relatif\nJarak dari pantai\nLaju penurunan tanah',PALE),(4.9,'Indeks kerentanan','Skor relatif tiap ruas\nTiga fitur berbobot sama',PALE),(9.05,'Skenario genangan','Indeks + pasut per jam\nEstimasi kedalaman',PANEL)]:
    box(s,x,2.42,3.62,2.20,col)
    text(s,title,x+.18,2.66,3.23,.49,25,LIGHT if col==PANEL else INK,True,DISPLAY)
    text(s,body,x+.18,3.34,3.2,1.22,20,SOFT if col==PANEL else MUTED)
text(s,'→',4.33,3.1,.54,.58,28,TEAL,True);text(s,'→',8.58,3.1,.54,.58,28,TEAL,True)
tag(s,'Pasut dari konstanta harmonik',.71,4.93,4.2)
text(s,'Pengamatan 2014; dibandingkan\ndengan IOC September 2026.',.71,5.45,4.32,.85,18,MUTED)
tag(s,'Aturan yang masih perlu diuji',5.02,4.93,4.3)
text(s,'Proporsi puncak 10% dan\nkedalaman 10 sampai 50 cm.',5.02,5.45,3.82,.85,18,MUTED)
tag(s,'Lalu mesin rute',9.13,4.93,3.5)
text(s,'Memakai kondisi pada\njam keberangkatan.',9.13,5.45,3.46,.85,18,MUTED)
text(s,'Umur data: OSM 2026 · subsidensi 2015 sampai 2018 · tahun akuisisi tile DEMNAS belum terverifikasi.',.7,6.55,11.93,.28,13.5,MUTED)
note(s,'Metode aktif',60,'Dzaky',
 'Ada dua jalur masukan. Kondisi wilayah membentuk indeks dari elevasi relatif, jarak pantai, dan laju penurunan tanah. Pasut dihitung per jam dari konstanta harmonik. Keduanya bertemu dalam aturan pemilihan ruas dan estimasi kedalaman, lalu dipakai mesin rute. Bobot tiga fitur sama besar. Proporsi puncak sepuluh persen dan kedalaman sepuluh sampai lima puluh sentimeter masih asumsi yang perlu diuji. Data subsidensi berasal dari 2015 sampai 2018, sedangkan konstanta pasut berasal dari 2014. Mesin rute memakai kondisi jam berangkat sepanjang perjalanan. Perubahan genangan ketika kendaraan sudah berjalan belum dimodelkan.',
 'Ikuti diagram dari kiri ke kanan. Jangan menyebut indeks sebagai probabilitas. Detail grid dan algoritme tersedia di lampiran 15.',
 'Audit data 25 September; script 11; backend/app/domain/{kerentanan,pasut,genangan,routing}.py.')

# 6. Dua jenis bukti tidak dicampur menjadi angka akurasi tunggal.
s=slide(6,'Bukti dan batas','Alur aplikasi sudah diuji. Genangan perlu data lapangan.',
 'Pengujian perangkat lunak dan pengujian pasut menjawab pertanyaan yang berbeda.',
 source='Audit regresi 22-23 September; uji pasut baru 18-25 September 2026. Rincian dan data tersedia di repo.',url=REPORT)
tag(s,'Perangkat lunak',.69,2.43,5.6)
text(s,'95 tes backend + 26 frontend',.69,2.93,6.0,.6,30,INK,True,DISPLAY)
text(s,'Retry, pergantian jam, data di luar cakupan,\nserta batas permintaan diuji.',.69,3.70,5.98,.86,21,MUTED)
text(s,'412 request uji VPS\n0 OOM / 0 restart selama 94,47 detik',.69,4.94,5.98,.98,24,INK,True,DISPLAY)
box(s,6.90,2.43,.016,3.45,LINE)
tag(s,'Pasut di stasiun',7.40,2.43,5.1)
text(s,'0,834',7.40,2.94,2.36,.73,39,TEAL,True,DISPLAY)
text(s,'korelasi',7.42,3.83,2.36,.4,19,MUTED)
text(s,'12,44 cm',10.03,2.94,2.72,.73,36,TEAL,True,DISPLAY)
text(s,'RMSE simpangan',10.03,3.83,2.72,.5,18,MUTED)
text(s,'10.020 rekaman, 18-25 September.\nModel dibekukan. Yang dibandingkan:\nsimpangan terhadap rata-rata deret.',7.42,4.69,5.24,1.16,18,MUTED)
box(s,.65,6.23,12.02,.56,PALE)
text(s,'Akurasi genangan per ruas dan hasil uji pengguna belum tersedia. Data IOC belum melalui pemeriksaan mutu.',.83,6.39,11.66,.28,15,INK)
note(s,'Apa yang sudah dibuktikan',70,'Dzaky',
 'Tes perangkat lunak memeriksa perilaku aplikasi: pencarian ulang, pergantian jam, dan penolakan waktu tanpa data. Pada audit VPS, 412 permintaan dijalankan selama sekitar satu setengah menit tanpa kehabisan memori atau restart. Ada respons sibuk yang memang dibatasi, jadi kami tidak menyebutnya semua permintaan berhasil.\n\nUntuk pasut, kami membandingkan model yang dibekukan dengan 10.020 rekaman IOC pada 18 sampai 25 September. Korelasinya 0,834 dan RMSE simpangannya 12,44 sentimeter. Itu pengujian pasut di stasiun. Akurasi genangan tiap ruas belum tersedia, dan data IOC sendiri belum melalui pemeriksaan mutu. Uji pengguna juga belum dilakukan. Sesudah ini, Naufal akan membahas keputusan yang bisa dibaca dari contoh perjalanan tadi.',
 'Jangan mengubah korelasi menjadi persen akurasi. Rincian sensor di slide 14 dan beban VPS di slide 16.',
 'audit_regresi_22_verifikasi.json; uji_pasut_independen_25_september.json; konfirmasi pengguna 25 September.')

# 7. Dampak sebagai keputusan dengan biaya, bukan penghematan fiktif.
s=slide(7,'Dampak | Smart Low-Carbon Urban Mobility','Memutar punya biaya. Jam berangkat ikut menentukan.',
 'Hasil model untuk motor, Tawang → Terboyo, 26 September 2026.',
 source='API potret lokal; dampak.py. Selisih dihitung sebelum pembulatan. Belum hasil perjalanan pengguna.')
grid_table(s,['Pilihan','Waktu','Jarak','Ruas basah model'],[
 ['08.00 | pembanding','7,0 mnt','6,32 km','29'],
 ['08.00 | sadar rob','13,9 mnt','10,58 km','14'],
 ['13.00 | sadar rob','7,0 mnt','6,32 km','0'],
],.66,2.28,[4.4,2.18,2.18,3.24],rowh=.69,size=20)
text(s,'Biaya memutar pukul 08.00',.72,5.18,4.75,.42,24,INK,True,DISPLAY)
text(s,'+6,9 menit  ·  +4,27 km',.72,5.84,5.35,.57,28,TEAL,True,DISPLAY)
text(s,'BBM +0,060 sampai 0,111 L\nCO₂ +0,138 sampai 0,256 kg',7.15,5.22,5.28,1.06,23,INK)
text(s,'Jika jadwal fleksibel, bandingkan jam lain. Biaya menunggu dan penghematan nyata belum diukur.',.73,6.57,11.9,.29,14.5,MUTED)
note(s,'Manfaat dan biaya adaptasi',60,'Naufal',
 'Pada jam delapan, rute sadar rob mengurangi jumlah ruas yang ditandai basah dalam model, dari 29 menjadi 14. Jaraknya justru bertambah 4,27 kilometer dan waktunya bertambah 6,9 menit. Artinya, menghindari sebagian paparan punya biaya. Estimasi bahan bakar dan CO₂ juga bertambah pada pilihan ini. Kalau jadwalnya fleksibel, pengguna bisa membandingkan jam satu siang. Menunggu lima jam tentu punya biaya jadwal yang belum kami hitung. Kaitan kami dengan Smart Low-Carbon Urban Mobility ada pada pembacaan konsekuensi energi dari pilihan perjalanan. Penghematan nyata baru bisa dinilai setelah ada perjalanan pengguna dan pembanding yang diukur.',
 'Tunjuk baris 08.00 dulu, lalu 13.00. Sebut CO₂ pembakaran, bukan CO₂e. Jangan menyebut rute 13.00 pasti kering.',
 'aset/skenario.json; backend/app/domain/dampak.py; subtema 4 rulebook halaman 3.')

# 8. Pengembangan berbasis bukti, dengan status yang jelas.
s=slide(8,'Langkah setelah final','Bawa prototipe ini ke uji pengguna dan ruas nyata.',
 'Rencana kerja berikut belum menjadi hasil atau kemitraan yang sudah berjalan.',
 source='Rencana tim berdasarkan audit 25 September. Target peserta merupakan rencana, bukan jumlah pengguna saat ini.')
items=[('Uji alur dengan 5 sampai 8 calon pengguna','Minta memilih tujuan, mengganti jam, dan menjelaskan peringatan.\nCatat tugas selesai, waktu, serta salah tafsir.'),
       ('Cocokkan hasil dengan pengamatan ruas','Kumpulkan lokasi, jam, kondisi basah atau kering, dan kedalaman.\nPisahkan kejadian untuk pengembangan dan pengujian.'),
       ('Jadwalkan pembaruan dan periksa sumbernya','Perbarui pasut, jalan, dan kondisi infrastruktur.\nPantau cakupan data sebelum habis.')]
for i,(title,body) in enumerate(items):
    yy=2.45+i*1.30
    text(s,f'{i+1:02}',.71,yy,.58,.52,27,TEAL,True,DISPLAY)
    text(s,title,1.52,yy,10.98,.47,25,INK,True,DISPLAY)
    text(s,body,1.53,yy+.58,10.9,.71,18.5,MUTED)
text(s,'Target tahap berikut: pengguna memahami batasnya, dan galat model dapat dihitung dari data lapangan.',.73,6.63,11.9,.26,14,MUTED)
note(s,'Rencana validasi berikut',55,'Naufal',
 'Langkah terdekat kami adalah uji tugas dengan lima sampai delapan calon pengguna. Yang dicari bukan sekadar pendapat bagus atau buruk. Kami ingin melihat apakah mereka bisa memilih tujuan, mengganti waktu, dan memahami bahwa kedalaman pada layar masih estimasi. Secara paralel, diperlukan pengamatan ruas dengan lokasi, jam, dan kedalaman yang jelas. Data itu dipisahkan untuk pengembangan dan pengujian. Pembaruan pasut serta data jalan juga perlu dijadwalkan. Aplikasi tidak bisa ditinggal dengan data beku selama tiga tahun. Kerja sama dengan pemilik data dan pengguna masih perlu dijajaki; belum ada mitra aktif yang kami klaim.',
 'Jangan menyebut target sebagai komitmen peserta. Bila waktu sudah 08.30, pakai dua kalimat: uji pengguna dan validasi ruas, lalu lanjut penutup.',
 'User mengonfirmasi belum ada uji pengguna; audit_asal_data_25_september.md.')

# 9. Penutup yang mengundang juri mencoba, bukan mengulang semua slide.
s=slide(9,'PASANG SURUT','Coba satu perjalanan. Bandingkan pilihan waktunya.',
 'Aplikasi dan kode dapat dibuka oleh juri.',dark=True,
 source='trio la albiceleste | Muhammad Dzaky Ahnaf · Daffa Rajendra Priyatama · Naufal Syafi\' Hakim')
qr=qrcode.make(APP);qr.save(ASSET/'qr_aplikasi.png')
box(s,.70,2.52,3.18,3.18,WHITE);picture(s,ASSET/'qr_aplikasi.png',.80,2.62,2.98)
text(s,'pasang-surut.vercel.app/app',4.55,2.75,7.96,.64,31,LIGHT,True,DISPLAY,url=APP)
text(s,'Perencanaan rute sadar kerentanan rob\nuntuk pesisir Semarang.',4.57,3.75,7.57,1.02,27,SOFT)
text(s,'Terima kasih, Bapak dan Ibu juri.',4.57,5.10,7.57,.59,25,LIGHT)
button(s,'Buka aplikasi',.76,6.25,2.10,url=APP,dark=True)
button(s,'Buka repositori',3.13,6.25,2.15,url=REPO,dark=True)
button(s,'Buka bukti data',5.55,6.25,2.20,url=REPORT,dark=True)
button(s,'Lampiran tanya jawab',9.96,6.25,2.58,target=10,dark=True)
note(s,'Penutup',25,'Naufal',
 'PASANG SURUT sudah bisa dipakai untuk mencoba satu perjalanan dan membandingkan pilihan waktunya. Pengujian genangan lapangan menjadi pekerjaan berikutnya. Aplikasi, kode, dan bukti pengujian kami buka melalui tautan ini. Terima kasih, Bapak dan Ibu juri. Kami siap membahas pertanyaan atau mencoba skenario lain bersama.',
 'Berhenti pada slide ini. Jangan lanjut otomatis ke lampiran. Bila juri ingin detail, klik menu lampiran atau ketik 10 lalu Enter.',
 'Alamat publik aplikasi dan repo.')

# 10. Menu lampiran, tombol kembali pada seluruh halaman sesudahnya.
s=slide(10,'Lampiran | dipakai saat ditanya','Bukti yang dapat dibuka saat tanya jawab.',
 'Klik topik, atau ketik nomor slide lalu Enter di PowerPoint.',source='Lampiran berada di luar sembilan menit presentasi utama.')
menu=[(11,'Demo 08.00'),(12,'Demo 13.00'),(13,'Sumber dan umur data'),(14,'Uji pasut September'),(15,'Indeks, genangan, dan rute'),(16,'Uji beban dan hosting'),(17,'Struktur kode dan reproduksi'),(18,'Perhitungan BBM dan CO₂')]
for i,(target,title) in enumerate(menu):
    x=.72+(i%2)*6.23;y=2.38+(i//2)*1.03
    button(s,f'{target:02}   {title}',x,y,5.84,target=target,size=18,height=.66)
button(s,'Kembali ke penutup',.74,6.61,2.5,target=9)
note(s,'Navigasi lampiran',0,'Sesuai topik',
 'Jawab singkat dahulu, lalu buka bukti yang diminta. Data dan metode ke Dzaky; demo dan antarmuka ke Daffa; dampak serta rencana uji ke Naufal. Anggota lain menambahkan setelah pembicara selesai.',
 'Satu pembicara menjawab satu pertanyaan. Jangan membuka semua lampiran secara berurutan.', 'Menu internal PPT.')

# 11 dan 12. Screenshot asli dengan hasil besar, semua tersimpan dalam PPTX.
for n,hour in [(11,'8'),(12,'13')]:
    props=next(f['properties'] for f in DATA[hour]['rute']['features'] if f['properties']['jenis']=='rute_sadar_rob')
    s=slide(n,'Cadangan demo',f'26 September 2026, pukul {int(hour):02}.00 WIB.',
      'Motor · akses Tawang → akses Terboyo · potret lokal',source='Screenshot aplikasi asli. Seluruh nilai pada halaman ini adalah keluaran model, bukan observasi jalan.')
    picture(s,OLD/f'demo_{int(hour):02}.png',.59,2.16,8.05)
    stat(s,dec(props['menit'])+' mnt','Waktu tempuh model',9.05,2.43,3.53)
    stat(s,dec(props['jarak_km'],2)+' km',str(props['ruas_tergenang'])+' ruas basah dalam model',9.05,4.11,3.53)
    text(s,'Masih menembus estimasi\ngenangan sampai 24,8 cm.' if hour=='8' else 'Nol pada model tidak\nmembuktikan jalan kering.',9.05,5.79,3.57,.86,18,RED if hour=='8' else MUTED,True)
    button(s,'Lanjut 13.00' if hour=='8' else 'Kembali ke metode',9.03,6.68,2.27,target=12 if hour=='8' else 5)
    note(s,f'Cadangan demo {hour}.00',0,'Daffa',
      'Ini tangkapan aplikasi dari potret yang sama dengan skenario presentasi. '+('Pada 08.00, hasilnya 13,9 menit, 10,58 kilometer dan 14 ruas basah dalam model. Kedalaman maksimum model 24,8 sentimeter, sehingga peringatan tetap perlu dibaca.' if hour=='8' else 'Pada 13.00, hasilnya 7 menit, 6,32 kilometer dan nol ruas basah dalam model. Hasil ini menjadi bahan perbandingan waktu, bukan bukti jalan sedang kering.'),
      'Gunakan bila live demo tertahan. Setelah halaman 12, kembali ke slide 5. Nomor 11 dan 12 dapat diketik langsung saat slideshow.',
      'docs/final/presentasi/aset/demo_08.png; demo_13.png; skenario.json.')

# 13. Sumber dan keberlanjutan dengan tanggal yang terbaca.
s=slide(13,'Data dan masa berlaku','Tanggal keluaran 2026 bukan tahun semua pengukuran.',
 'Database aktif memakai kerentanan_v1. Tidak ada sensor genangan langsung per ruas.',source='Audit asal data 25 September 2026; metadata berkas dan respons API produksi.',url=REPORT)
grid_table(s,['Masukan','Asal','Periode / status'],[
 ['Jalan dan tempat','OpenStreetMap','Graf 24 Agustus 2026'],
 ['Konstanta pasut','Rachman dkk.','Pengamatan 13-27 Maret 2014'],
 ['Elevasi','DEMNAS BIG','Tahun akuisisi tile belum terverifikasi'],
 ['Subsidensi','Rahmawati dkk.','Citra 2015 sampai 2018'],
],.67,2.24,[3.03,3.15,5.82],rowh=.63,size=18)
text(s,'Data produksi berakhir 5 Oktober 2026, 07.00 WIB.',.75,5.57,11.8,.46,25,INK,True,DISPLAY)
text(s,'Potret lokal sampai 29 September, 00.00 WIB. Pipeline pembaruan permanen belum ada.\nOperasi jangka panjang memerlukan pembaruan, validasi, pemantauan, dan penanggung jawab data.',.75,6.17,11.8,.64,17,MUTED)
note(s,'Sumber dan keberlanjutan',0,'Dzaky',
 'Sumbernya dapat ditelusuri, tetapi umurnya berbeda. Jalan diekstrak pada 2026, konstanta pasut berasal dari 2014, dan subsidensi memakai periode 2015 sampai 2018 pada tingkat kecamatan. Tahun akuisisi raster belum terverifikasi. Pasut sudah dibandingkan dengan rekaman September, tetapi genangan per ruas belum. Untuk tiga tahun operasi, kami perlu data yang diperbarui dan diuji berkala. Saat audit, produksi berakhir 5 Oktober 07.00 WIB; jadwal permanen penerbit data belum tersedia. IOC juga membatasi penggunaan komersial; rencana produk perlu pembicaraan dengan pemilik data.',
 'Bila juri meminta data mentah, buka tautan audit. Jangan mengekstrapolasi laju subsidensi lama secara linear lalu menyebutnya elevasi 2029.', 'audit_asal_data_25_september.md; data/referensi; status_produksi_25_september.json.')

# 14. Bukti sensor baru, bukan angka akurasi genangan.
s=slide(14,'Validasi pasut','Uji September memakai model yang dibekukan.',
 'Tidak memilih ulang offset atau amplitudo berdasarkan hasil minggu ini.',source='IOC/VLIZ sema, sensor prs; 18-25 September 2026. CSV dan skrip pengujian tersedia pada laporan audit.',url=REPORT)
picture(s,OUT.parent/'data/perbandingan_pasut_september.png',.96,2.20,11.4)
note(s,'Uji pasut independen',0,'Dzaky',
 'Ada 10.020 rekaman sensor tekanan sema. Model lama dan offset tujuh jam dipertahankan. Korelasi 0,8340, RMSE simpangan 12,44 sentimeter, dan MAE 10,32 sentimeter. Titik nol alat dan rekonstruksi berbeda; setiap deret dikurangi rata-rata periode sebelum dibandingkan. Sensor pada layanan IOC belum melalui quality control, dan sampel yang berdekatan saling berkorelasi. Satu minggu ini tidak membuktikan ketelitian lintas musim, ketepatan tinggi absolut, atau akurasi genangan ruas. Uji Agustus memakai jendela yang ikut memilih offset, sehingga sifatnya berbeda dengan uji baru ini.',
 'Tunjuk selisih kurva di puncak. Jangan menyatakan r 0,834 sebagai akurasi 83,4%.', 'data/uji_pasut_independen_25_september.json; CSV IOC; disclaimer IOC.')

# 15. Algoritme dan keputusan menghentikan model ML.
s=slide(15,'Metode dan batas','Indeks, kedalaman, dan rute punya asumsi yang berbeda.',
 'Eksperimen Sentinel-1 tidak dipakai sebagai prediktor rilis.',source='kerentanan.py; genangan.py; routing.py; metrik_model.json; audit data 25 September.',url=REPO+'/tree/main/backend/app/domain')
grid_table(s,['Tahap','Yang dilakukan','Batas saat ini'],[
 ['Indeks','Elevasi relatif, pantai, subsidensi; bobot ⅓','Skor relatif, bukan peluang banjir'],
 ['Lingkungan elevasi','Median 3 × 3 sel, sisi sel 500 m','Bukan radius lingkaran 500 m'],
 ['Genangan','Pasut mengatur proporsi; kedalaman 10-50 cm','Parameter belum terkalibrasi lapangan'],
 ['Rute','Bobot tetap pada jam keberangkatan','Belum mengikuti perubahan saat perjalanan'],
],.66,2.30,[1.85,5.44,4.71],rowh=.75,size=16.5)
text(s,'Model Sentinel-1 ditolak: ROC-AUC 0,6579 · PR-AUC 0,0371 · F1 0,0894.',.75,6.19,11.8,.5,22,INK,True,DISPLAY)
text(s,'Angka model eksperimen tersebut tidak menjadi metrik akurasi indeks yang aktif.',.75,6.73,11.8,.25,14,MUTED)
note(s,'Model dan routing',0,'Dzaky',
 'Indeks memakai tiga fitur sama besar, dinormalisasi dalam wilayah studi. Elevasi relatif memakai sembilan sel grid, bukan radius lingkaran. Indeks dan pasut lalu diubah menjadi estimasi genangan dengan parameter yang belum dikalibrasi. Hujan, debit sungai, pompa dan tanggul belum masuk rumus aktif. Model Sentinel-1 pernah dilatih, tetapi label dan hasilnya tidak layak dijadikan prediktor rilis. Routing sekarang menggunakan kondisi jam berangkat yang tetap agar konsisten. Uji 2.400 pencarian pada 40 graf kecil dibandingkan dengan Bellman-Ford menguji model bobot tetap itu; tidak membuktikan optimalitas pada genangan yang berubah selama perjalanan.',
 'Bila ditanya kenapa bukan AI, jelaskan keputusan berdasarkan hasil eksperimen. Jangan menyebut indeks sebagai model ML terlatih.', 'backend/tests/test_routing_properti.py; metrik_model.json; domain routing dan kerentanan.')

# 16. Ketahanan terukur dan arsitektur operasional yang aktual.
s=slide(16,'Implementasi dan ketahanan','Permintaan berat dibatasi; masa berlaku data diperiksa.',
 'Browser → Vercel → VPS PASANG SURUT (frontend, API, PostgreSQL/PostGIS).',source='Audit VPS 22 September: 412 request / 94,47 detik. Tidak mengukur kapasitas maksimum atau SLA.',url=REPO+'/blob/main/docs/final/audit_regresi_22_september.md')
text(s,'0 OOM / 0 restart',.7,2.48,6.15,.69,36,TEAL,True,DISPLAY)
text(s,'316 respons 200\n32 respons sibuk terkendali\n64 penolakan masukan ekstrem',.73,3.49,5.9,1.62,23,INK)
text(s,'Puncak sampel API 186,89 MiB\ndengan batas container 320 MiB.',.73,5.64,5.9,.84,20,MUTED)
for i,(title,body) in enumerate([('Saat permintaan bertumpuk','Dua request berat bersamaan; selebihnya\ndapat menerima respons sibuk.'),('Saat waktu tidak tercakup','Permintaan ditolak. Tidak diterjemahkan\nmenjadi jalan kering.'),('Saat internet acara bermasalah','Demo lokal memakai potret yang sama\ndengan screenshot cadangan.')]):
    yy=2.44+i*1.39
    text(s,title,7.24,yy,5.36,.48,24,INK,True,DISPLAY)
    text(s,body,7.24,yy+.57,5.26,.76,18,MUTED)
note(s,'Hosting dan uji beban',0,'Dzaky / Daffa',
 'Render sebelumnya restart karena melewati 512 MB. Perbaikan mencakup penggunaan bersama graf, pembatasan cache dan request berat, lalu migrasi layanan ke VPS yang diuji bertahap. Dalam audit tambahan, 412 permintaan selesai dengan status yang diharapkan: 316 berhasil, 32 sibuk, 58 payload terlalu besar, empat waktu atau koordinat di luar cakupan, dan dua masukan buruk. Tidak ada OOM atau restart dalam jendela 94,47 detik. Ini bukan janji kapasitas tanpa batas. Origin HTTP 18080 masih sementara; jalur origin terenkripsi dan penguatan operasional masuk pekerjaan setelah final. Layanan lain di VPS tidak diubah.',
 'Jika ditanya kecepatan maksimum atau SLA, jawab belum diukur. Angka audit memori berbeda dari audit sebelumnya karena kondisi cache berbeda.', 'audit_regresi_22_requests.json; audit_regresi_22_verifikasi.json; deploy/vps/README.md.')

# 17. Bukti kualitas kode yang dapat diklik.
s=slide(17,'Code project','Kode dipisah supaya perhitungan bisa diuji sendiri.',
 'Tautan repositori mengarah ke bagian yang ditunjukkan.',source='Source code publik, tes regresi, CI, serta dokumentasi deploy dan pemulihan.',url=REPO)
for x,y,title,path,body in [(.74,2.49,'Antarmuka','frontend/src','Pilihan tujuan, pergantian jam,\nretry, status data dan peta.'),(7.08,2.49,'API dan cakupan','backend/app','Validasi masukan, versi data,\ncache dan penanganan galat.'),(.74,4.52,'Perhitungan domain','backend/app/domain','Pasut, indeks, genangan, rute,\nserta estimasi dampak.'),(7.08,4.52,'Data dan pengujian','backend/tests','Uji regresi, graf pembanding,\nkasus ekstrem dan cakupan waktu.')]:
    rule(s,x,y,5.49,TEAL);text(s,title,x,y+.25,5.44,.49,26,INK,True,DISPLAY)
    text(s,body,x,y+.91,5.44,.82,20,MUTED)
    text(s,path,x,y+1.70,5.4,.25,13,TEAL,True,url=REPO+'/tree/main/'+path)
note(s,'Struktur kode dan reproduksi',0,'Dzaky / Daffa',
 'Perhitungan berada di domain sehingga tidak harus diuji lewat tampilan peta. API memeriksa waktu, koordinat dan versi data. Frontend menangani permintaan yang dibatalkan saat jam berubah serta pencarian ulang. Backend memiliki tes regresi dan pembanding algoritme; CI menjalankan tes pada lingkungan bersih. Data dan script pipeline terpisah dari image API. Untuk reproduksi demo, gunakan deploy/demo_lokal.py dan konfigurasi Vite khusus demo. Panduan backup dan pemulihan database ada pada deploy/vps. Repositori terbuka bagi juri.',
 'Klik direktori yang ditanyakan saja. Jangan melakukan deploy, migrasi, atau tes beban baru di depan juri.', 'README; backend/app; backend/tests; frontend/src/App.test.jsx; .github/workflows/uji.yml; deploy/vps.')

# 18. Akuntansi dampak dan batas penafsiran.
s=slide(18,'Perhitungan dampak','BBM dan CO₂ dihitung dari tambahan jarak.',
 'Pembanding mengabaikan genangan pada jam dan moda yang sama.',source='dampak.py; ambang_moda; skenario.json. Rationale faktor di docs/sumber_angka.md; konsumsi masih asumsi.',url=REPO+'/blob/main/docs/sumber_angka.md')
text(s,'Δ BBM = Δ jarak × konsumsi per km\nΔ CO₂ = Δ BBM × faktor pembakaran',.73,2.40,11.9,1.12,31,INK,True,DISPLAY)
rule(s,.73,3.89,11.89)
stat(s,'0,020 L/km','Asumsi konsumsi motor',.76,4.28,3.6)
stat(s,'±30%','Rentang asumsi konsumsi',4.99,4.28,3.6)
stat(s,'2,31 kg/L','Faktor CO₂ pembakaran bensin',9.16,4.28,3.5)
text(s,'Belum mengukur kemacetan nyata, biaya menunggu, penghematan kota, atau dampak kesehatan pengguna.',.77,6.44,11.83,.48,17,MUTED)
note(s,'Asumsi BBM dan emisi',0,'Naufal / Dzaky',
 'Selisih jarak dikalikan konsumsi wakil moda, lalu faktor pembakaran. Untuk motor digunakan 0,020 liter per kilometer dengan rentang sensitivitas tiga puluh persen, dan faktor 2,31 kilogram CO₂ per liter bensin. Ada penjelasan literatur faktor di docs/sumber_angka.md, tetapi konsumsi kendaraan dan rentang sensitivitas tetap asumsi rilis, bukan hasil pengukuran perjalanan ini. CO₂ tersebut tidak mencakup gas lain atau daur hidup. Kami belum punya dasar untuk mengubah hasil satu skenario menjadi penghematan satu kota atau dampak kesehatan terukur.',
 'Gunakan tanda tambah ketika membicarakan skenario 08.00. Jangan menyebut biaya menghindar sebagai penghematan.', 'backend/app/domain/dampak.py; docs/sumber_angka.md; db/schema.sql; aset/skenario.json.')

# Navigasi internal dan lampiran yang dilewati saat F5.
for i,s in enumerate(slides[10:],start=11):
    button(s,'Menu',11.94,.30,.77,target=10)
for shape,target in links:shape.click_action.target_slide=slides[target-1]
for s in slides[9:]:s._element.set('show','0')
for s in slides:
    for sh in s.shapes:
        style=sh._element.find(qn('p:style'))
        if style is not None:
            ef=style.find(qn('a:effectRef'))
            if ef is not None:ef.set('idx','0')
        sp=sh._element.find(qn('p:spPr'))
        if sp is not None and sp.find(qn('a:effectLst')) is None:sp.append(OxmlElement('a:effectLst'))
assert sum(n['seconds'] for n in notes)==540
# Sembilan menit memberi cadangan 60 detik terhadap batas sepuluh menit.
P.save(OUT/f'{NAME}.pptx')
(OUT/'catatan_pembicara.md').write_text('# Catatan pembicara\n\nPembagian latihan; target 9 menit. Belum hasil latihan tim dengan stopwatch.\n\n'+'\n\n---\n\n'.join(f"## {n['slide']:02}. {n['title']}\n\n{n['time']} | {n['speaker']}\n\n{n['body']}\n\nAksi operator: {n['cues']}\n\nSumber: {n['sources']}" for n in notes)+'\n',encoding='utf-8')
manifest={'slides':len(slides),'main_slides':9,'appendices':9,'target_seconds':540,'buffer_seconds':60,
 'fonts':[FONT,DISPLAY],'application_url':APP,'scenario_date':'2026-09-26','scenario_version':DATA['8']['versi_data'],
 'source_code_base':'69eae60','user_testing':'Belum dilakukan; dikonfirmasi user 25 September 2026',
 'references':['Rulebook-dsdc-final.pdf halaman 2,3,9,10','TM Final DSDC Anforcom 2026.pdf halaman 5,6,7,9',
               'ITC2026_1_SOFTDEV_PPT_Fable5Enjoyer_DekapAutis_Revisi_v3.pptx dan PDF','anti-slop-writing/indonesian/SKILL.md'],
 'slides_summary':[{k:v for k,v in n.items() if k not in ['body','cues','sources']} for n in notes],
 'main_narration_words':sum(len(n['body'].split()) for n in notes[:9]),
 'assets_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [OLD/'demo_08.png',OLD/'demo_13.png',OLD/'skenario.json',OUT.parent/'data/perbandingan_pasut_september.png']}}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k not in ['slides_summary','assets_sha256','references']},ensure_ascii=True,indent=2))
