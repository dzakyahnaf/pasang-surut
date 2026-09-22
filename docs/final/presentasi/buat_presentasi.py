"""Bangun deck final dari bukti repo; dependency khusus materi, bukan aplikasi.

Jalankan dari akar repo. Paket dapat diinstal ke .deploy-local/presentation-tools.
PowerPoint mengekspor PDF melalui ekspor_presentasi.ps1. Semua teks, tabel,
diagram, dan peta konteks adalah objek PowerPoint yang bisa diedit.
"""
from __future__ import annotations
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / '.deploy-local/presentation-tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
import qrcode

OUT = Path(__file__).resolve().parent
ASSET = OUT / 'aset'
P = Presentation()
P.slide_width = Inches(13.333333)
P.slide_height = Inches(7.5)
P.core_properties.title = 'PASANG SURUT | Final ANFORCOM 2026'
P.core_properties.subject = 'Perencanaan perjalanan sadar rob di pesisir Semarang'
P.core_properties.author = 'trio la albiceleste — ITS Surabaya'
P.core_properties.keywords = 'ANFORCOM 2026, Semarang, pasut, routing, kerentanan'
P.core_properties.comments = '8 slide utama, 6 lampiran. Disusun 23 September 2026 dari bukti dalam repo.'
NAVY='0B1F2A'; PANEL='12303E'; LINE='CBD7DB'; BG='F2F5F6'; WHITE='FFFFFF'
INK='0B1F2A'; MUTED='46626F'; TEAL='1C7F9E'; CYAN='4FB3C4'; PALE='E3EAEC'
AMBER='FFB020'; LIGHT='E8F1F4'; RED='C8322B'
FONT='Arial'; DISPLAY='Arial Narrow'
REPO='https://github.com/dzakyahnaf/pasang-surut'
APP='https://pasang-surut.vercel.app/app'
DATA=json.loads((ASSET/'skenario.json').read_text(encoding='utf-8'))
TIDE=json.loads((ROOT/'data/referensi/kalibrasi_pasut.json').read_text(encoding='utf-8'))
AUDIT=json.loads((ROOT/'docs/final/bukti/audit_regresi_22_verifikasi.json').read_text(encoding='utf-8'))
notes=[]

def rgb(value): return RGBColor.from_string(value)
def box(s,x,y,w,h,color,stroke=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                          Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color)
    if stroke: sh.line.color.rgb=rgb(stroke); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    return sh

def text(s,value,x,y,w,h,size=22,color=INK,bold=False,font=FONT,align=None,link=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(5); p.space_before=Pt(0); p.line_spacing=1.04
        if align is not None: p.alignment=align
        r=p.add_run(); r.text=line
        r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
        r.font.language_id=1057
    if link: sh.click_action.hyperlink.address=link
    return sh

def rule(s,x,y,w,color=LINE): return box(s,x,y,w,.015,color)
def label(s,t,x,y,w=8,color=TEAL): return text(s,t.upper(),x,y,w,.3,12,color,True)
def footer(s,n,source='',dark=False,url=None):
    color='ABC0CA' if dark else MUTED
    rule(s,.55,7.05,12.23,PANEL if dark else LINE)
    text(s,source or 'trio la albiceleste  /  ITS Surabaya  /  ANFORCOM 2026',.55,7.15,11.7,.2,10,color,link=url)
    text(s,f'{n:02d}',12.35,7.11,.42,.27,14,color,True,align=PP_ALIGN.RIGHT)

def slide(n,chapter,title,subtitle=None,dark=False,source='',url=None):
    s=P.slides.add_slide(P.slide_layouts[6]); s.background.fill.solid()
    s.background.fill.fore_color.rgb=rgb(NAVY if dark else BG)
    label(s,chapter,.55,.35,10,CYAN if dark else TEAL)
    text(s,title,.55,.85,12.25,.63,34,LIGHT if dark else INK,True,DISPLAY)
    if subtitle: text(s,subtitle,.57,1.56,12.0,.52,17,'ABC0CA' if dark else MUTED)
    footer(s,n,source,dark,url)
    return s

def note(s,title,timing,speaker,body,sources=''):
    value=f'{title}\nWAKTU: {timing}\nPEMBICARA (usulan): {speaker}\n\n{body}\n\nSUMBER: {sources}'
    s.notes_slide.notes_text_frame.text=value
    notes.append(value)

def metric(s,value,caption,x,y,w=3,color=TEAL,detail=None):
    text(s,value,x,y,w,.7,38,color,True,DISPLAY)
    text(s,caption,x,y+.86,w,.65,20,INK,True)
    if detail: text(s,detail,x,y+1.57,w,.75,16,MUTED)

def image(s,name,x,y,w,h=None):
    return s.shapes.add_picture(str(ASSET/name),Inches(x),Inches(y),width=Inches(w),height=Inches(h) if h else None)

def poly(s,paths,color,width=1):
    builder=None
    for path in paths:
        if len(path)<2: continue
        pts=[(round(x*1000),round(y*1000)) for x,y in path]
        if builder is None: builder=s.shapes.build_freeform(*pts[0],scale=Inches(1)/1000)
        else: builder.move_to(*pts[0])
        builder.add_line_segments(pts[1:],close=False)
    if builder:
        sh=builder.convert_to_shape(); sh.fill.background(); sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
        return sh

def context_map(s,x,y,w,h):
    # Peta vektor dari ekstrak OSM tersimpan. Hanya ruas utama untuk orientasi.
    bbox=(110.405,-6.978,110.486,-6.935)
    def inside(p): return bbox[0]<=p[0]<=bbox[2] and bbox[1]<=p[1]<=bbox[3]
    def xy(p): return (x+(p[0]-bbox[0])/(bbox[2]-bbox[0])*w,y+(bbox[3]-p[1])/(bbox[3]-bbox[1])*h)
    box(s,x,y,w,h,PALE)
    roads=json.loads((ROOT/'data/processed/ruas_jalan.geojson').read_text(encoding='utf-8'))
    paths=[]
    for f in roads['features']:
        if f['properties'].get('jenis') not in ['trunk','primary','secondary','tertiary','trunk_link','primary_link']: continue
        coords=f['geometry']['coordinates']
        part=[]
        for p in coords:
            if inside(p): part.append(xy(p))
            else:
                if len(part)>1: paths.append(part)
                part=[]
        if len(part)>1: paths.append(part)
    poly(s,paths,'A0B5BF',.7)
    coast=json.loads((ROOT/'data/processed/garis_pantai.geojson').read_text(encoding='utf-8'))
    paths=[]
    for f in coast['features']:
        seqs=[f['geometry']['coordinates']] if f['geometry']['type']=='LineString' else f['geometry']['coordinates']
        for seq in seqs:
            part=[]
            for p in seq:
                if inside(p): part.append(xy(p))
                else:
                    if len(part)>1: paths.append(part)
                    part=[]
            if len(part)>1: paths.append(part)
    poly(s,paths,TEAL,1.5)
    text(s,'LAUT JAWA',x+.18,y+.15,2,.3,13,TEAL,True)
    for coords,title,dx,dy in [([110.42744,-6.96374],'TAWANG',-.32,.16),([110.46934,-6.94825],'TERBOYO',-.56,-.44),([110.41853,-6.94959],'TANJUNG EMAS',-.58,-.42)]:
        px,py=xy(coords)
        dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(px-.055),Inches(py-.055),Inches(.11),Inches(.11))
        dot.fill.solid(); dot.fill.fore_color.rgb=rgb(INK); dot.line.fill.background()
        text(s,title,px+dx,py+dy,1.9,.28,13,INK,True)
    text(s,'SEMARANG',x+2.0,y+h-.48,2.8,.35,18,TEAL,True)
    text(s,'U ↑',x+w-.48,y+.15,.42,.3,13,INK)

# 01 — pertanyaan pengguna dan konteks tempat.
s=P.slides.add_slide(P.slide_layouts[6]); s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY)
label(s,'Final DSDC / ANFORCOM 2026',.55,.4,7,CYAN)
text(s,'PASANG\nSURUT',.55,1.02,5.7,2.07,58,LIGHT,True,DISPLAY)
text(s,'Berangkat kapan,\nlewat mana?',.6,3.26,5.4,1.35,36,LIGHT,True,DISPLAY)
text(s,'Perencanaan perjalanan sadar rob\ndi pesisir Semarang.',.6,4.8,5.1,.92,23,'ABC0CA')
box(s,6.36,1.02,6.43,4.86,PALE)
context_map(s,6.36,1.5,6.43,3.44)
text(s,'Contoh perjalanan: akses Tawang → akses Terboyo',6.4,6.05,6.2,.4,16,LIGHT)
text(s,'trio la albiceleste  •  ITS Surabaya',.6,5.9,5.45,.33,17,LIGHT,True)
text(s,"Muhammad Dzaky Ahnaf  /  Daffa Rajendra Priyatama\nNaufal Syafi’ Hakim",.6,6.34,5.5,.53,12,'ABC0CA')
footer(s,1,'Peta konteks: ekstrak OpenStreetMap tersimpan; penanda menunjukkan titik akses.',True,'https://www.openstreetmap.org/copyright')
note(s,'01 — Masalah pengguna','00.00–00.45 (45 detik)','Dzaky',
'''Bayangkan seorang pengendara yang perlu berangkat dari sekitar Tawang menuju kawasan industri Terboyo. Pada kawasan pesisir, pilihannya tidak berhenti pada “jalan mana yang paling cepat”. Ia juga perlu mempertimbangkan jam berangkat dan kondisi ruas yang dilalui.
PASANG SURUT membantu membandingkan pilihan perjalanan itu. Kami fokus pada pesisir Semarang, dengan tujuan agar waktu, rute, dan konsekuensi memutar dapat dibaca dalam satu alur.
Skenario ini ilustrasi kebutuhan pengguna, bukan kutipan wawancara atau bukti aplikasi sudah dipakai warga. Penanda pada peta adalah akses jaringan jalan, bukan pintu bangunan yang telah diverifikasi.''',
'data/processed/tujuan_cepat.geojson; ruas_jalan.geojson; garis_pantai.geojson; proposal tim.')

# 02 — keputusan pengguna, bukan daftar teknologi.
s=slide(2,'01 / Masalah pengguna → keputusan','Satu perjalanan, tiga informasi',
        'Pengendara di pesisir Semarang membandingkan pilihan sebelum berangkat.')
cols=[(.6,'01','Pilih perjalanan','Asal–tujuan bernama\ndan moda kendaraan.'),
      (4.84,'02','Bandingkan waktu','Geser jam dan lihat\nperubahan pilihan rute.'),
      (9.08,'03','Baca konsekuensi','Waktu, jarak, paparan,\nserta estimasi BBM/CO₂.')]
for x,n,title,body in cols:
    rule(s,x,2.5,3.65,TEAL)
    text(s,n,x,2.78,1,.6,30,TEAL,True,DISPLAY)
    text(s,title,x,3.65,3.6,.5,25,INK,True,DISPLAY)
    text(s,body,x,4.38,3.65,1.25,22,MUTED)
box(s,.6,6.08,12.1,.6,PALE)
text(s,'Indeks kerentanan memberi konteks perencanaan. Kondisi lapangan tetap perlu diperiksa.',.82,6.25,11.7,.28,16,INK)
note(s,'02 — Nilai produk','00.45–01.30 (45 detik)','Dzaky',
'''Ada tiga informasi yang kami satukan. Pertama, asal–tujuan dan kendaraan, karena batas motor berbeda dengan mobil. Kedua, waktu keberangkatan yang dapat digeser untuk membandingkan kondisi model. Ketiga, konsekuensi rute yang dipilih: tambahan waktu, jarak, serta estimasi biaya bahan bakar dan emisi.
Sumber yang aktif adalah indeks kerentanan; kedalaman yang terlihat bukan hasil pengukuran langsung. Karena itu aplikasi juga menampilkan sumber, waktu data, dan peringatan paparan.
Daffa akan memperlihatkan satu perjalanan yang sama pada dua jam berbeda.''',
'frontend/src/components/PanelRute.jsx, PitaPasut.jsx, PanelDampak.jsx; copy.id.json.')

# 03 — ruang demo, gambar cadangan tetap ada.
s=slide(3,'02 / Demo','Tawang → Terboyo, pada dua waktu',
        'Motor • potret 26 September 2026 • 08.00 dan 13.00 WIB • sumber kerentanan_v1',
        source='Skenario tersimpan 919aa345…; hasil API lengkap: aset/skenario.json')
image(s,'demo_08.png',.6,2.16,8.0,4.8)
for yy,num,title,detail in [(2.3,'1','Kenali lokasi','Tawang dan Terboyo\nterlihat pada peta.'),(3.63,'2','Baca peringatan','08.00: memutar, tetapi\nmasih ada paparan.'),(4.96,'3','Geser jam','13.00: tanpa tambahan\njarak menurut model.')]:
    text(s,num,9.2,yy,.45,.44,23,TEAL,True)
    text(s,title,9.78,yy,2.8,.42,23,INK,True,DISPLAY)
    text(s,detail,9.78,yy+.51,2.95,.79,18,MUTED)
text(s,'Buka aplikasi ↗',9.2,6.45,3.3,.38,19,TEAL,True,link=APP)
note(s,'03 — Demo','01.30–04.30 (3 menit, termasuk perpindahan)','Daffa',
'''00–30 detik: buka aplikasi yang sudah siap. Sebut Semarang, asal akses Tawang, tujuan akses Terboyo, moda motor, serta tanggal/sumber yang terlihat. Gambar di slide berasal dari potret lokal 26 September, bukan keadaan lapangan hari ini.
30–90 detik: pilih 26 September 08.00 WIB. Tunjukkan rute sadar rob 13,9 menit/10,58 km dan pembanding 7,0 menit/6,32 km yang mengabaikan genangan. Maksimum model pada rute sadar rob 24,8 cm; rute masih menembus ruas berisiko. Jangan sebut rute ini “aman”.
90–135 detik: geser ke 13.00 WIB. Tunggu pembaruan selesai. Pada potret ini rute sadar rob dan pembanding sama: 7,0 menit/6,32 km. Tidak ada genangan yang dimodelkan pada rute; ini tidak membuktikan jalan benar-benar kering.
135–165 detik: simpulkan pilihan yang dapat dipertimbangkan bila jadwal fleksibel, sambil tetap memeriksa kondisi lapangan. Waktu 7,0 menit adalah model kecepatan, bukan ETA lalu lintas langsung.
165–180 detik: kembali ke slide metode.
Jika produksi tertahan lebih dari sekitar 10 detik, gunakan lokal yang sudah menyala. Jika lokal gagal, buka slide 9 lalu 10 (nomor + Enter); kembali ke slide 4. Screenshot sudah tertanam. MP4 belum termasuk paket ini. Jangan menunggu server berulang kali.
Produksi memakai dataset berbeda dari potret; jangan menjanjikan angka persis sama. Untuk cerita yang konsisten, demo lokal potret adalah pilihan siap pakai; cek cakupan produksi sebelum tampil.''',
'aset/demo_08.png; aset/demo_13.png; aset/skenario.json; data/processed/tujuan_cepat.geojson.')

# 04 — jalur data dipisahkan dari runtime.
s=slide(4,'03 / Metode','Dari pasut dan kerentanan ke pilihan rute',
        '19.394 ruas pada wilayah pilot; data disiapkan sebelum pengguna mencari perjalanan.',
        source='Sumber implementasi: backend/app/domain/{pasut,kerentanan,routing}.py; runtime.py',url=REPO+'/tree/main/backend/app')
for yy,t,b in [(2.27,'Rekonstruksi pasut','Komponen waktu berdasarkan konstanta'),(3.66,'Indeks kerentanan','Elevasi relatif, pantai, dan subsidensi')]:
    box(s,.6,yy,4.23,1.13,PANEL)
    text(s,t,.82,yy+.15,3.82,.39,24,LIGHT,True,DISPLAY)
    text(s,b,.82,yy+.69,3.84,.25,15,'ABC0CA')
poly(s,[[(4.91,2.82),(5.3,2.82),(5.3,4.2),(4.91,4.2)],[(5.3,3.5),(5.67,3.5)]],TEAL,1.8)
text(s,'→',5.56,3.25,.5,.5,29,TEAL,True)
for x,t,b in [(6.18,'Kondisi per jam','Estimasi untuk ruas\ndengan batas data'),(10.06,'Dua rute','Sadar rob dan\npembanding')]:
    box(s,x,2.47,2.67,2.28,PANEL)
    text(s,t,x+.18,2.7,2.31,.85,24,LIGHT,True,DISPLAY)
    text(s,b,x+.18,3.77,2.31,.82,17,'ABC0CA')
text(s,'→',9.18,3.15,.5,.5,29,TEAL,True)
box(s,.6,5.07,12.13,.65,PALE)
text(s,'Saat klik: browser/peta → API → database atau potret → hasil perbandingan',.83,5.27,11.66,.35,20,INK,True,DISPLAY)
text(s,'Kondisi jam keberangkatan dipakai sepanjang rute. Perubahan selama perjalanan belum dimodelkan.',.65,6.02,12.0,.78,21,MUTED)
note(s,'04 — Metode','04.30–05.30 (60 detik)','Dzaky',
'''Pasut memberi komponen waktu. Sementara itu, indeks kerentanan membedakan ruas berdasarkan elevasi relatif, jarak dari pantai, dan laju subsidensi. Ketiga komponen memakai bobot sama sebagai asumsi, belum hasil kalibrasi genangan lapangan.
Kondisi per jam disiapkan dalam dataset, lalu API membandingkan rute sadar rob dan rute yang mengabaikan genangan. Klik pengguna tidak mengunduh ulang graf OSM atau melatih model. Browser membaca hasil dari API; dataset dapat berasal dari database VPS atau potret untuk demo lokal.
Untuk menghindari masalah pergantian jam pada algoritme lama, rilis final memakai kondisi jam keberangkatan sepanjang satu pencarian. Kami tidak mengklaim rute optimal terhadap kondisi yang berubah selama perjalanan.
Jika ditanya elevasi: median dari kumpulan 3 × 3 sel, sisi sel 500 meter. Ini bukan radius melingkar 500 meter.''',
'backend/app/domain/kerentanan.py; backend/app/domain/routing.py; backend/app/runtime.py. OSM sebagai graf; DEMNAS sebagai fitur relatif.')

# 05 — dua jenis bukti dan batasnya.
s=slide(5,'04 / Bukti pengujian dan batasan','Pisahkan bukti perangkat lunak dan model',
        source='Audit 22 Sep + regresi panel 23 Sep; kalibrasi IOC 28 Agu. Detail dan sumber pada lampiran 11–13.')
label(s,'Perangkat lunak',.6,2.0,5)
metric(s,'95 / 26','Tes backend / frontend',.6,2.55,5.25,detail='Backend: Windows & Linux. Frontend: 26 tes\nsetelah koreksi panel pada 23 September.')
rule(s,6.34,2.0,.015,LINE);box(s,6.34,2.0,.015,3.95,LINE)
label(s,'Rekonstruksi pasut',6.72,2.0,5.7)
metric(s,'r = 0,782','RMSE 0,116 m',6.72,2.55,5.75,detail='Jendela kalibrasi 10 hari; simpangan muka air.\nEvaluasi ini belum merupakan uji independen.')
box(s,.6,5.56,12.1,1.07,PANEL)
text(s,'Belum ada validasi genangan per ruas dan dampak perjalanan di lapangan.',.86,5.78,11.52,.44,24,LIGHT,True,DISPLAY)
text(s,'Lulus regresi tidak berarti seluruh bug hilang. Akurasi pasut tidak menjadi akurasi genangan.',.86,6.29,11.5,.25,14,'ABC0CA')
note(s,'05 — Bukti dan batasan','05.30–06.30 (60 detik)','Dzaky',
'''Ada dua lapis bukti yang perlu dibedakan. Pada perangkat lunak, 95 tes backend lulus di Windows dan container Linux. Frontend kini memiliki 26 tes setelah peringatan kedalaman di panel rute dikoreksi. Salah satu tes backend membandingkan 2.400 pencarian pada graf kecil dengan Bellman-Ford; angka itu bagian dari tes, bukan tambahan 2.400 unit test.
Pada pasut, pilihan acuan WIB memberi korelasi 0,7821 dan RMSE 0,1155 meter dalam jendela sepuluh hari terhadap rekaman IOC. Perbandingan memakai simpangan muka air karena datum berbeda. Data tersebut juga digunakan untuk memilih acuan waktu, sehingga hasilnya bukan evaluasi independen pada data baru.
Kami belum punya pengamatan genangan per ruas untuk menguji indeks, dan belum punya bukti dampak perjalanan lapangan. Karena itu kami melaporkan kedalaman dan biaya sebagai estimasi serta menampilkan batasnya pada antarmuka.''',
'docs/final/audit_regresi_22_september.md; frontend/src/App.test.jsx; data/referensi/kalibrasi_pasut.json (hipotesis_wib); docs/validasi.md §3.1.')

# 06 — tabel hasil, angka dari respons aktual, bukan proyeksi dampak kota.
s=slide(6,'05 / Dampak','Memahami biaya adaptasi perjalanan',
        'Simulasi motor • akses Tawang → akses Terboyo • potret 26 September 2026',
        source='Angka dihitung ulang dari API lokal rilis final: aset/skenario.json. Faktor BBM/CO₂ masih asumsi.')
def route(hour,kind): return next(f['properties'] for f in DATA[str(hour)]['rute']['features'] if f['properties']['jenis']==kind)
def dec(v,n=1): return f'{v:.{n}f}'.replace('.',',')
r8=route(8,'rute_sadar_rob'); rb=route(8,'rute_abai_rob');r13=route(13,'rute_sadar_rob');d8=DATA['8']['dampak']
xs=[.6,4.65,7.35,10.05]; ws=[4.05,2.7,2.7,2.65]
heads=['Pilihan pada potret','Pembanding\n08.00 WIB','Sadar rob\n08.00 WIB','Sadar rob\n13.00 WIB']
for x,w,h in zip(xs,ws,heads):
    box(s,x,2.2,w,.87,PANEL);text(s,h,x+.17,2.29,w-.32,.74,19,LIGHT,True,DISPLAY)
rows=[('Waktu model (menit)',*[dec(r['menit']) for r in [rb,r8,r13]]),
      ('Jarak (km)',*[dec(r['jarak_km'],2) for r in [rb,r8,r13]]),
      ('Ruas tergenang menurut model',*[str(r['ruas_tergenang']) for r in [rb,r8,r13]])]
for i,row in enumerate(rows):
    yy=3.07+i*.61
    for j,(x,w,v) in enumerate(zip(xs,ws,row)):
        box(s,x,yy,w,.61,PALE if i%2==0 else BG)
        text(s,v,x+.17,yy+.13,w-.32,.38,17 if j==0 else 22,INK,j!=0)
box(s,.6,5.23,12.1,1.48,PALE)
text(s,f'08.00: +{dec(d8["menit"])} menit  /  +{dec(d8["km"],2)} km untuk menghindar',.82,5.44,11.65,.42,25,INK,True,DISPLAY)
text(s,f'Estimasi tambahan BBM {dec(d8["liter"]["bawah"],3)}–{dec(d8["liter"]["atas"],3)} L  •  CO₂ {dec(d8["kg_co2"]["bawah"],3)}–{dec(d8["kg_co2"]["atas"],3)} kg',.82,5.98,11.65,.33,19,TEAL,True)
text(s,'Rute 08.00 masih memiliki paparan. Hasil 13.00 tidak membuktikan jalan benar-benar kering.',.82,6.42,11.65,.24,14,MUTED)
note(s,'06 — Dampak yang dapat dijelaskan','06.30–07.30 (60 detik)','Naufal',
'''Kami menyebut angka ini biaya adaptasi. Pembanding pukul delapan mengabaikan genangan: 7,0 menit dan 6,32 kilometer. Rute sadar rob memerlukan 13,9 menit dan 10,58 kilometer. Model menunjukkan jumlah ruas tergenang yang dilalui turun dari 29 menjadi 14, tetapi rute masih memiliki paparan. Jumlah ruas bukan pengukuran risiko atau lama paparan.
Memutar menambah sekitar 6,9 menit dan 4,27 kilometer. Dari asumsi konsumsi motor, itu setara tambahan 0,060 sampai 0,111 liter BBM serta 0,138 sampai 0,256 kilogram CO₂ pembakaran. Ini rentang asumsi, bukan interval kepercayaan dan bukan penurunan emisi yang sudah diukur.
Pada pukul tiga belas di potret ini, kedua rute sama dan model tidak memberi genangan di jalur. Bila jadwal fleksibel, informasi ini dapat menjadi bahan mempertimbangkan jam lain. Menunggu lima jam juga memiliki biaya jadwal yang belum kami hitung. Jangan menyebut selisih ini penghematan total atau dampak kota.''',
'aset/skenario.json, motor, versi 919aa345-7492-4a53-9e7c-c34086775d18; backend/app/domain/dampak.py.')

# 07 — tahap selanjutnya dengan keluaran yang bisa diperiksa.
s=slide(7,'06 / Rencana pengembangan','Dari prototipe menuju pilot yang terukur',
        'Usulan tahap berikutnya; belum diklaim sebagai kegiatan atau kemitraan yang selesai.')
stages=[(.6,'01 / KEGUNAAN','Uji tugas 5–8 peserta','Apakah lokasi, jam, dan\nperingatan dipahami?','Ukur: keberhasilan tugas\ndan waktu penyelesaian.'),
        (4.84,'02 / MODEL','Observasi ruas bertanggal','Catat lokasi, waktu,\nkedalaman, dan kondisi jalan.','Uji pada data terpisah;\nlaporkan galat dan cakupan.'),
        (9.08,'03 / OPERASI','Pembaruan & pemantauan','Tetapkan pemilik dataset,\njadwal publikasi, dan biaya.','Pantau umur data, latensi,\nRAM, serta galat layanan.')]
for x,k,t,b,m in stages:
    label(s,k,x,2.53,3.66)
    text(s,t,x,3.11,3.65,.93,27,INK,True,DISPLAY)
    text(s,b,x,4.25,3.65,1.1,19,MUTED)
    rule(s,x,5.54,3.65,LINE)
    text(s,m,x,5.84,3.65,.75,18,TEAL)
note(s,'07 — Rencana pengembangan','07.30–08.30 (60 detik)','Naufal',
'''Tahap selanjutnya dimulai dengan uji tugas pada lima sampai delapan peserta yang relevan, misalnya pengendara atau pengiriman di wilayah pilot. Kami ingin mengukur apakah mereka dapat memilih asal–tujuan, membaca waktu, serta memahami bahwa peringatan adalah estimasi. Jumlah ini rencana, bukan peserta yang sudah diuji.
Berikutnya, kami memerlukan observasi ruas yang memiliki lokasi dan waktu jelas. Data baru dipisahkan untuk evaluasi, sehingga bobot dan estimasi tidak dinilai pada data yang sama dengan penyetelannya. Calon instansi atau komunitas belum disebut sebagai mitra aktif.
Secara operasional, kami perlu jadwal publikasi data, penanggung jawab pemeliharaan, pemantauan memori serta latensi, dan biaya hosting. Rilis saat ini belum punya scheduler publikasi otomatis. Perluasan wilayah dilakukan setelah bukti pada pilot membaik.''',
'docs/final/outline_slide.md; docs/final/audit_regresi_22_september.md; rencana tim, belum hasil uji pengguna.')

# 08 — penutup, QR lokal, tautan dapat diklik.
s=slide(8,'PASANG SURUT / trio la albiceleste','Bandingkan rute. Pilih waktu.',
        dark=True,source='ANFORCOM 2026  •  Smart Low-Carbon Urban Mobility  •  ITS Surabaya')
text(s,'Pahami\nkonsekuensinya.',.6,2.25,7.55,1.9,52,LIGHT,True,DISPLAY)
text(s,'Informasi perjalanan pesisir Semarang\nyang bisa dibandingkan, dengan batas\nbukti yang terlihat.',.65,4.45,7.32,1.5,25,'ABC0CA')
qr=qrcode.QRCode(version=None,error_correction=qrcode.constants.ERROR_CORRECT_M,box_size=14,border=4)
qr.add_data(APP);qr.make(fit=True);qr.make_image(fill_color='#0B1F2A',back_color='white').save(ASSET/'qr_aplikasi.png')
image(s,'qr_aplikasi.png',9.27,2.32,2.76,2.76)
text(s,'COBA APLIKASI',9.0,5.4,3.33,.35,18,LIGHT,True,align=PP_ALIGN.CENTER,link=APP)
text(s,'pasang-surut.vercel.app/app',8.55,5.96,4.2,.32,15,'ABC0CA',align=PP_ALIGN.CENTER,link=APP)
text(s,'Kode & bukti pengujian ↗',.66,6.28,7,.34,18,CYAN,True,link=REPO)
note(s,'08 — Penutup','08.30–09.00 (30 detik)','Naufal',
'''PASANG SURUT membantu pengguna membandingkan rute, mempertimbangkan waktu, dan memahami konsekuensi perjalanan di pesisir Semarang. Produk yang kami tunjukkan sudah memiliki alur interaktif dan bukti pengujian perangkat lunak. Pengujian genangan dan dampak lapangan adalah langkah berikutnya.
Aplikasi dan kode tersedia melalui tautan ini. Kami siap menjelaskan implementasi serta batas bukti yang kami punya. Terima kasih.
Berhenti pada slide ini. Slide 9–14 adalah lampiran, bukan lanjutan presentasi utama. Untuk Q&A, ketik nomor slide lalu Enter. Pembagian pembicara adalah usulan; lakukan latihan bersama sebelum dibekukan.''',
APP+'; '+REPO)

# 09–10 — screenshot tertanam: tampilkan dengan nomor slide saat demo gagal.
for n,h in [(9,8),(10,13)]:
    rr=route(h,'rute_sadar_rob')
    s=slide(n,'Lampiran / Cadangan demo',f'26 September, {h:02d}.00 WIB — motor',
            source='Gambar aplikasi lokal, 23 Sep 2026. Potret bertanggal; bukan laporan kondisi lapangan.')
    image(s,f'demo_{h:02d}.png',.6,1.78,8.55,5.13)
    text(s,f'{dec(rr["menit"])}',9.55,2.05,3,.74,44,TEAL,True,DISPLAY)
    text(s,'menit menurut model',9.55,2.88,3,.38,18,MUTED)
    text(s,f'{dec(rr["jarak_km"],2)} km',9.55,3.5,3,.62,32,INK,True,DISPLAY)
    text(s,'Masih menembus\n14 ruas tergenang\nmenurut model.' if h==8 else 'Sama dengan pembanding.\nModel tidak memberi\ngenangan di rute.',9.55,4.45,3.1,1.52,20,INK)
    text(s,'Kembali ke metode:\ntekan 4 lalu Enter',9.55,6.13,3.15,.64,16,TEAL)
    note(s,f'{n:02d} — Cadangan demo {h:02d}.00','Di dalam alokasi 3 menit demo bila diperlukan','Daffa',
         f'Gunakan narasi slide 3 untuk jam {h:02d}.00. Gambar asli tertanam; tidak memerlukan jaringan untuk ditampilkan. Sumber kerentanan_v1, versi potret 919aa345…; mode motor. Jangan menyebut angka ini observasi lapangan. Kembali ke slide 4 setelah kedua kondisi ditunjukkan.',
         f'aset/demo_{h:02d}.png; aset/skenario.json.')

# 11 — bukti pasut, tidak mencampur offset terbaik dengan offset yang dipakai.
s=slide(11,'Lampiran / Pasut','Yang diuji: simpangan muka air',
        'Kalibrasi 28 Agustus 2026 • IOC stasiun sema • acuan fase yang dipakai: WIB (+7 jam)',
        source='docs/validasi.md §3.1; data/referensi/kalibrasi_pasut.json. Batang adalah korelasi pada WIB.',url=REPO+'/blob/main/docs/validasi.md')
text(s,'Korelasi menurut panjang jendela',.6,2.22,6.2,.4,22,INK,True,DISPLAY)
for i,(days,val) in enumerate([(2,.907),(4,.909),(7,.858),(10,.7821)]):
    yy=3.02+i*.69
    text(s,f'{days} hari',.65,yy,1.05,.32,18,MUTED)
    box(s,1.88,yy+.015,4.1,.32,PALE)
    box(s,1.88,yy+.015,4.1*val,.32,TEAL)
    text(s,dec(val,3),6.14,yy,1.1,.36,20,INK,True)
box(s,7.83,2.2,4.9,3.02,PANEL)
text(s,'10 hari / 12.317 rekaman',8.06,2.45,4.45,.48,24,LIGHT,True,DISPLAY)
text(s,'r = 0,7821\nRMSE = 0,1155 m',8.06,3.23,4.2,1.1,30,LIGHT,True,DISPLAY)
text(s,'Datum berbeda: bandingkan\nsimpangan terhadap rata-rata.',8.06,4.45,4.32,.7,17,'ABC0CA')
text(s,'Data juga dipakai memilih acuan waktu. Perlu evaluasi pada jendela baru;\nkoreksi nodal belum diterapkan. Hasil ini tidak mengukur genangan jalan.',.65,5.88,12.0,.83,22,MUTED)
note(s,'11 — Q&A pasut','Lampiran, di luar 9 menit','Dzaky',
'''Nilai 0,7821 dan 0,1155 m berasal dari hipotesis WIB pada jendela 10 hari. Jangan tertukar dengan offset terbaik +8 jam yang menghasilkan r 0,8465/RMSE 0,097 m: sistem memilih +7 jam, bukan offset yang paling cocok pada sampel.
Jendela 2/4/7 hari pada docs/validasi.md memberi korelasi WIB 0,907/0,909/0,858. Jendela saling tumpang tindih, bukan empat uji independen. Rekaman IOC untuk 10 hari berjumlah 12.317. Karena datum berbeda, yang dibandingkan simpangan terhadap rata-rata.
Acuan dipilih menggunakan data yang sama; evaluasi ini bersifat kalibrasi/in-sample. Konstanta berasal dari rekaman 15 hari pada 2014, P1/K2 diturunkan dari komponen lain, dan koreksi nodal belum diterapkan. Belum ada klaim akurasi genangan per ruas.''',
'data/referensi/kalibrasi_pasut.json; docs/validasi.md; backend/scripts/04_kalibrasi_pasut.py.')

# 12 — indeks operasional dan eksperimen yang ditolak.
s=slide(12,'Lampiran / Pemilihan metode','Mengapa rilis memakai indeks kerentanan?',
        source='kerentanan.py; metrik_model.json; koreksi metode pada docs/final/audit_regresi_22_september.md')
label(s,'Yang dipakai pada rilis',.6,2.0,5.75)
text(s,'3 fitur, bobot sama',.6,2.6,5.7,.52,29,INK,True,DISPLAY)
text(s,'Elevasi relatif • jarak pantai • subsidensi.\nBobot ⅓ adalah asumsi, belum terkalibrasi.',.6,3.24,5.8,.78,20,MUTED)
for row in range(3):
    for col in range(3): box(s,.63+col*.5,4.52+row*.5,.45,.45,TEAL if row==col==1 else LINE)
text(s,'Median 3 × 3 sel\nSisi sel 500 m\nBukan radius melingkar',2.42,4.54,3.78,1.3,21,INK)
box(s,6.65,2.0,.015,4.52,LINE)
label(s,'Eksperimen tidak dipakai',7.07,2.0,5.66)
text(s,'Model berlabel Sentinel-1',7.07,2.6,5.67,.52,28,INK,True,DISPLAY)
text(s,'Latih 2015–2023; uji 2024–2026.\nROC-AUC 0,6579 • PR-AUC 0,0371\nF1 0,0894 • proporsi dasar PR 0,016',7.07,3.24,5.63,1.29,21,MUTED)
box(s,7.07,4.93,5.62,1.56,PALE)
text(s,'Label basah satelit belum menjadi\nbukti genangan jalan. Model ini\ntidak dipakai sebagai prediktor rilis.',7.28,5.12,5.17,1.25,21,INK)
note(s,'12 — Q&A indeks dan Sentinel-1','Lampiran, di luar 9 menit','Dzaky',
'''Indeks memakai tiga skor relatif, dinormalisasi, dengan bobot sepertiga. Elevasi relatif adalah elevasi ruas dikurangi median dari anggota 3 × 3 sel grid dengan sisi 500 m. Nama metadata lama radius_elevasi_relatif_m dipertahankan pada artefak lama, tetapi bukan jendela Euclidean radius 500 m. Tinggi DEM tidak dipotong langsung dengan muka air untuk membuat genangan.
Eksperimen HistGradientBoostingClassifier dengan label perubahan VV Sentinel-1 tidak dipakai pada rilis. Angka uji berdasarkan split waktu tercatat apa adanya: ROC-AUC 0,6579, PR-AUC 0,0371, F1 0,0894. Kelas basah pada uji 1,6%; metrik klasifikasi label satelit tidak mengukur akurasi genangan jalan. Pengujian lanjutan label dan kaitan muka air tidak menghasilkan landasan yang cukup untuk memakai model sebagai prediksi rob.
Keputusan rilis adalah menampilkan indeks dengan batas klaim jelas. Kami masih memerlukan observasi genangan per ruas pada waktu tertentu untuk mengevaluasinya.''',
'backend/app/domain/kerentanan.py; data/processed/indeks_kerentanan.json; data/referensi/metrik_model.json; uji_muka_air_terukur.json; docs/validasi.md.')

# 13 — kapasitas dan regresi, hasil uji terikat durasi/lingkungan.
s=slide(13,'Lampiran / Keandalan','Regresi dan uji VPS setelah perbaikan',
        'Audit 22 September 2026 • API satu worker • 412 request dalam 94,47 detik',
        source='audit_regresi_22_verifikasi.json; audit_regresi_22_requests.json. Bukan uji kapasitas maksimum.',url=REPO+'/blob/main/docs/final/audit_regresi_22_september.md')
text(s,'Puncak sampel RAM / batas container',.6,2.23,6.3,.44,24,INK,True,DISPLAY)
for i,(k,cap,name) in enumerate([('/pasang-surut-api-1',320,'API'),('/pasang-surut-db-1',192,'DB'),('/pasang-surut-web-1',64,'Web')]):
    peak=AUDIT['telemetry']['max_usage_mib'][k]; yy=3.03+i*.85
    text(s,name,.6,yy,1.05,.38,20,INK,True)
    box(s,1.76,yy+.02,3.21,.33,PALE);box(s,1.76,yy+.02,3.21*peak/cap,.33,TEAL)
    text(s,f'{dec(peak,1)} / {cap} MiB',5.2,yy,2.35,.4,18,MUTED)
box(s,8.0,2.25,4.7,3.68,PANEL)
text(s,'0 OOM / 0 restart',8.22,2.56,4.25,.53,28,LIGHT,True,DISPLAY)
text(s,'32 respons sibuk terkendali.\n316 berhasil; sisanya masukan\nyang memang harus ditolak.',8.22,3.43,4.17,1.22,20,'ABC0CA')
text(s,'RAM host tersedia minimum\n645,92 MiB pada sampel uji.',8.22,5.0,4.1,.75,18,LIGHT)
text(s,'2.400 pencarian dibanding Bellman-Ford: kondisi tetap pada jam keberangkatan.\nData rusak, waktu di luar cakupan, retry, batas body, serta versi campuran diuji.',.65,6.13,12.0,.65,18,MUTED)
note(s,'13 — Q&A pengujian dan hosting','Lampiran, di luar 9 menit','Dzaky / Daffa',
'''Insiden Render memakai lebih dari 512 MB dan restart. Rilis sekarang memakai VPS bersama, dengan isolasi container API 320 MiB, database 192 MiB, web 64 MiB, satu worker API, cache graf bersama, dan pembatasan dua request berat bersamaan.
Audit tambahan 22 September berisi 412 request dalam 94,47 detik: 316 HTTP200, 58 HTTP413, empat HTTP422, dua HTTP400, dan 32 HTTP503 sibuk yang diharapkan. Tidak ada respons tak terduga, OOM, restart, atau kenaikan failcnt. Puncak dari 188 sampel: API186,89; DB93,20; web51,71 MiB. Minimum RAM host tersedia645,92 MiB. Uji memakai guard sumber daya; bukan DDoS atau kapasitas maksimum.
Uji beban terdahulu 1.596 request selama7m41d juga tersedia di laporan terpisah; jangan menjumlahkan atau menyamakan puncak memori karena keadaan cache berbeda. Layanan Maknaprice tidak diubah; identitas container/config sama dan semua10 probe tambahan HTTP200.
95 tes backend mencakup satu tes properti dengan2.400 perbandingan pada40 graf kecil, bukan bukti optimalitas dinamis. Frontend26 tes per23 September. Batas20 detik request dan10 detik SQL bukan jaminan kecepatan. Browser fisik/proyektor tetap harus diuji tim.
Vercel adalah alamat masuk; runtime API/frontend/database ada pada VPS. Origin HTTP18080 bersifat sementara dengan izin pemilik. Kode audit main1651ad3; tag APIaudit22-11c58776b1ca. Koreksi copy panel23 September tidak mengubah algoritme backend.''',
'docs/final/audit_regresi_22_september.md; docs/final/uji_beban_vps_22_september.md; backend/tests/test_routing_properti.py; frontend/src/App.test.jsx.')

# 14 — akuntansi, asumsi belum bersitasi diberi nama apa adanya.
s=slide(14,'Lampiran / Batas angka dampak','Angka BBM dan CO₂ adalah estimasi asumsi',
        'Selisih dibanding rute yang mengabaikan genangan pada jam dan moda yang sama.',
        source='backend/app/domain/dampak.py; aset/skenario.json; docs/batasan.md. Sitasi faktor belum dilengkapi.')
text(s,'Δ BBM = Δ jarak × konsumsi per km\nΔ CO₂ = Δ BBM × faktor pembakaran',.65,2.31,11.9,1.05,30,INK,True,DISPLAY)
rule(s,.65,3.71,12.02,LINE)
for x,val,cap in [(.65,'0,020 L/km','Konsumsi motor\nAsumsi rilis'),(4.87,'±30%','Rentang konsumsi\nAsumsi, bukan interval kepercayaan'),(9.09,'2,31 kg/L','Faktor CO₂ pembakaran\nSitasi belum dilengkapi')]:
    text(s,val,x,4.09,3.56,.63,33,TEAL,True,DISPLAY)
    text(s,cap,x,4.94,3.55,1.16,18,MUTED)
box(s,.6,6.2,12.1,.51,PALE)
text(s,'Belum mencakup kemacetan nyata, biaya menunggu, emisi daur hidup, atau manfaat kesehatan terukur.',.82,6.36,11.65,.26,15,INK)
note(s,'14 — Q&A dampak','Lampiran, di luar 9 menit','Naufal / Dzaky',
'''Rumus dihitung atas selisih jarak dua rute. Nilai konsumsi motor0,020 L/km, rentang±30%, dan faktor2,31 kg CO₂/L dibaca dari konfigurasi rilis. Rentang tersebut adalah asumsi sensitivitas konsumsi; bukan rentang hasil pengukuran atau interval kepercayaan.
Repositori belum melengkapi sitasi faktor; data/referensi/faktor_emisi.json masih template null. Karena itu kami menyebutnya asumsi implementasi dan tidak mengklaim metodologi inventarisasi emisi yang tervalidasi. Jika ditanya dampak kota, jawab belum tersedia. Perlu data pemakaian, perjalanan nyata, kendaraan, baseline, dan biaya perubahan jadwal.
Istilah CO₂e tidak dipakai: faktor ini hanya pembakaranCO₂, tidak memasukkan gas lain atau siklus hidup. Ketika salah satu rute tidak tersedia, API mengembalikan dampak null, bukan “hemat nol”.''',
'backend/app/domain/dampak.py; db/schema.sql; docs/batasan.md; data/referensi/faktor_emisi.json; aset/skenario.json.')

# Lampiran tersembunyi di slideshow biasa; PDF ekspor menyertakannya.
for s in list(P.slides)[8:]: s._element.set('show','0')
# Hindari efek bayangan default theme pada bentuk dan garis.
for s in P.slides:
    for sh in s.shapes:
        style=sh._element.find(qn('p:style'))
        if style is not None:
            effect=style.find(qn('a:effectRef'))
            if effect is not None: effect.set('idx','0')
        sppr=sh._element.find(qn('p:spPr'))
        if sppr is not None and sppr.find(qn('a:effectLst')) is None:
            sppr.append(OxmlElement('a:effectLst'))
P.save(OUT/'Pasang_Surut_Final_ANFORCOM_2026.pptx')
(OUT/'catatan_pembicara.md').write_text('# Catatan pembicara PASANG SURUT\n\nUsulan pembagian; target 9 menit, maksimum 10 menit.\n\n'+'\n\n---\n\n'.join(notes)+'\n',encoding='utf-8')
manifest={'slides':len(P.slides),'main_slides':8,'appendices':6,'target_seconds':540,
          'scenario_date':'2026-09-26','scenario_version':DATA['8']['versi_data'],
          'application_url':APP,'fonts':[FONT,DISPLAY],'source_code_base':'1651ad3',
          'frontend_panel_fix_commit':'c554e27',
          'frontend_tests_after_panel_fix':26,'backend_tests_audit_22':95}
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print(json.dumps(manifest))
